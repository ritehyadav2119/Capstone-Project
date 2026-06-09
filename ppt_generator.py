import os

from pptx import Presentation


class PPTGenerator:

    def __init__(self):

        os.makedirs(
            "exports/ppt",
            exist_ok=True
        )

    # ------------------------
    # ADD TITLE SLIDE
    # ------------------------

    def add_title_slide(
            self,
            prs,
            title,
            subtitle):

        slide = prs.slides.add_slide(
            prs.slide_layouts[0]
        )

        slide.shapes.title.text = (
            title
        )

        slide.placeholders[1].text = (
            subtitle
        )

    # ------------------------
    # ADD CONTENT SLIDE
    # ------------------------

    def add_content_slide(
            self,
            prs,
            title,
            content):

        slide = prs.slides.add_slide(
            prs.slide_layouts[1]
        )

        slide.shapes.title.text = (
            title
        )

        slide.placeholders[1].text = (
            content[:4000]
        )

    # ------------------------
    # GENERATE PPT
    # ------------------------

    def generate_presentation(
            self,
            document_name,
            summary,
            keywords,
            conclusion):

        prs = Presentation()

        self.add_title_slide(

            prs,

            "Intelligent Document Summarization System",

            document_name
        )

        self.add_content_slide(

            prs,

            "Document Summary",

            summary
        )

        self.add_content_slide(

            prs,

            "Keywords",

            keywords
        )

        self.add_content_slide(

            prs,

            "Conclusion",

            conclusion
        )

        self.add_content_slide(

            prs,

            "Future Scope",

            """
• AI Quiz Generation

• Flashcards

• Multi-language Support

• Voice Assistant

• Cloud Deployment
"""
        )

        output_path = (
            f"exports/ppt/"
            f"{document_name}.pptx"
        )

        prs.save(
            output_path
        )

        return output_path

    # ------------------------
    # AUTO PPT
    # ------------------------

    def generate_from_summary_pack(
            self,
            document_name,
            summary_pack):

        return self.generate_presentation(

            document_name,

            summary_pack.get(
                "short_summary",
                ""
            ),

            summary_pack.get(
                "keywords",
                ""
            ),

            summary_pack.get(
                "executive_summary",
                ""
            )
        )