import os

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)

from reportlab.lib.styles import (
    getSampleStyleSheet
)

from docx import Document


class ExportManager:

    def __init__(self):

        os.makedirs(
            "exports/txt",
            exist_ok=True
        )

        os.makedirs(
            "exports/pdf",
            exist_ok=True
        )

        os.makedirs(
            "exports/docx",
            exist_ok=True
        )

    # --------------------------
    # TXT EXPORT
    # --------------------------

    def export_txt(
            self,
            content,
            filename=
            "summary.txt"):

        path = os.path.join(

            "exports",

            "txt",

            filename
        )

        with open(

            path,

            "w",

            encoding="utf-8"

        ) as f:

            f.write(content)

        return path

    # --------------------------
    # PDF EXPORT
    # --------------------------

    def export_pdf(
            self,
            title,
            content,
            filename=
            "summary.pdf"):

        path = os.path.join(

            "exports",

            "pdf",

            filename
        )

        doc = (
            SimpleDocTemplate(
                path
            )
        )

        styles = (
            getSampleStyleSheet()
        )

        story = []

        story.append(

            Paragraph(
                title,
                styles["Title"]
            )
        )

        story.append(
            Spacer(
                1,
                12
            )
        )

        story.append(

            Paragraph(

                content.replace(
                    "\n",
                    "<br/>"
                ),

                styles["BodyText"]
            )
        )

        doc.build(
            story
        )

        return path

    # --------------------------
    # DOCX EXPORT
    # --------------------------

    def export_docx(
            self,
            title,
            content,
            filename=
            "summary.docx"):

        path = os.path.join(

            "exports",

            "docx",

            filename
        )

        doc = Document()

        doc.add_heading(
            title,
            level=1
        )

        doc.add_paragraph(
            content
        )

        doc.save(
            path
        )

        return path

    # --------------------------
    # CHAT HISTORY
    # --------------------------

    def export_chat_history(
            self,
            history):

        content = ""

        for item in history:

            content += (
                f"\nUser: "
                f"{item['question']}\n"
            )

            content += (
                f"Assistant: "
                f"{item['answer']}\n"
            )

            content += (
                "-" * 50 +
                "\n"
            )

        return self.export_txt(

            content,

            "chat_history.txt"
        )

    # --------------------------
    # QUIZ EXPORT
    # --------------------------

    def export_quiz(
            self,
            quiz_text):

        return self.export_pdf(

            "AI Generated Quiz",

            quiz_text,

            "quiz.pdf"
        )

    # --------------------------
    # FLASHCARDS
    # --------------------------

    def export_flashcards(
            self,
            flashcards):

        return self.export_docx(

            "Flashcards",

            flashcards,

            "flashcards.docx"
        )

    # --------------------------
    # COMPLETE REPORT
    # --------------------------

    def export_complete_report(
            self,
            summary,
            analytics,
            filename=
            "complete_report.pdf"):

        report = f"""

DOCUMENT SUMMARY

{summary}


DOCUMENT ANALYTICS

{analytics}

"""

        return self.export_pdf(

            "Document Report",

            report,

            filename
        )